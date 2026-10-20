"""最後逐筆人工對照與簽核。未完成前禁止進入正式計算。"""
from pathlib import Path
from datetime import datetime, date, timezone
import csv
import hashlib
import json
import os
import shutil
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from pypdf import PdfReader

FIELDS = ("test_no", "cycle1_mbar", "cycle1_pa", "cycle2_mbar",
          "cycle2_pa", "comments", "instrument_result")
HEADERS = ["Page", "Time", "Test No.", "Cycle 1 mbar", "Cycle 1 Pa",
           "Cycle 2 mbar", "Cycle 2 Pa", "Comments", "Result",
           "Disposition", "Source Image", "Confirmed", "Notes"]


class PendingReview(ValueError):
    pass


def source_rows(folder):
    rows = []
    with (folder / "CCIT.PDF").open("rb") as file:
        pages = len(PdfReader(file).pages)
    for page in range(1, pages + 1):
        with (folder / "ocr_output" / f"page_{page}_paired_columns.csv").open(
                encoding="utf-8-sig", newline="") as file:
            records = list(csv.DictReader(file))
        if not records:
            raise ValueError(f"第 {page} 頁沒有資料")
        for record in records:
            values = [int(record["test_no"])]
            values += [float(record[key]) for key in FIELDS[1:5]]
            values += [record["comments"], record["instrument_result"]]
            rows.append([page, record["time"], *values])
    for index, row in enumerate(rows):
        disposition = ("Exclude: last 10 suitability" if index >= len(rows) - 10
                       else "Exclude: QC Check" if "qc check" in row[7].casefold()
                       else "Keep")
        rows[index] += [disposition, f"ocr_output/page_{row[0]}.png"]
    return rows


def fingerprint(folder, rows):
    payload = {"pdf": hashlib.sha256((folder / "CCIT.PDF").read_bytes()).hexdigest(),
               "product_type": os.environ.get("CCIT_PRODUCT_TYPE", ""),
               "rows": rows, "review_version": 1}
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False,
                                    sort_keys=True).encode("utf-8")).hexdigest()


def create_review(path, rows, digest):
    book = Workbook()
    sign = book.active
    sign.title = "Sign-off"
    for row in [
        ["Final OCR Review", "Pending manual review"],
        ["Instructions", "Compare every row with the source image. Set Confirmed to Yes only after checking."],
        ["If incorrect", "Leave Confirmed blank. Correct upstream OCR review data, then rerun. Do not edit source values here."],
        ["Record count", len(rows)],
        ["Kept records", sum(r[9] == "Keep" for r in rows)],
        ["Excluded records", sum(r[9] != "Keep" for r in rows)],
        ["Reviewer name", None],
        ["Review date (YYYY-MM-DD)", None],
        ["Approval (Approved)", None],
        ["Declaration", "I checked all rows, quantities, names, results, total count and proposed exclusions against the source PDF."],
        ["Record type", "Manual sign-off record; no identity authentication or certified electronic signature."],
        ["Source ID", digest],
    ]:
        sign.append(row)
    sign.column_dimensions["A"].width = 31
    sign.column_dimensions["B"].width = 100
    for row in sign:
        row[1].alignment = Alignment(wrap_text=True, vertical="center")
        sign.row_dimensions[row[0].row].height = 34
    for cell in ("B7", "B8", "B9"):
        sign[cell].protection = Protection(locked=False)
        sign[cell].fill = PatternFill("solid", fgColor="FFF2CC")
    sign["B8"].number_format = "yyyy-mm-dd"
    approval = DataValidation(type="list", formula1='"Approved"', allow_blank=True)
    approval.errorTitle, approval.error = "Invalid approval", "Select Approved after reviewing all rows."
    approval.showErrorMessage = True
    sign.add_data_validation(approval)
    approval.add(sign["B9"])
    sign.protection.sheet = True
    sheet = book.create_sheet("Data Review")
    sheet.append(HEADERS)
    for values in rows:
        sheet.append(values + [None, None])
        n = sheet.max_row
        sheet.cell(n, 11).hyperlink = values[10]
        sheet.cell(n, 11).style = "Hyperlink"
        for col in (12, 13):
            sheet.cell(n, col).protection = Protection(locked=False)
            sheet.cell(n, col).fill = PatternFill("solid", fgColor="FFF2CC")
        for col in range(4, 8):
            sheet.cell(n, col).number_format = "0.0#######"
        if values[9] != "Keep":
            for cell in sheet[n][:10]:
                cell.fill = PatternFill("solid", fgColor="E7E6E6")
    confirmed = DataValidation(type="list", formula1='"Yes"', allow_blank=True)
    confirmed.showErrorMessage = True
    sheet.add_data_validation(confirmed)
    confirmed.add(f"L2:L{sheet.max_row}")
    for col, width in zip("ABCDEFGHIJKLM", (9, 14, 12, 17, 16, 17, 16, 44, 12, 35, 32, 15, 40)):
        sheet.column_dimensions[col].width = width
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="203B5D")
        cell.alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.protection.sheet = True
    sheet.protection.autoFilter = False
    book.save(path)


def verify_review(path, rows, digest):
    book = load_workbook(path, data_only=False)
    sheet, sign = book["Data Review"], book["Sign-off"]
    if sign["B12"].value != digest or sheet.max_row != len(rows) + 1:
        raise PendingReview("核對表來源或筆數不符，請重新產生核對表。")
    if [cell.value for cell in sheet[1]] != HEADERS:
        raise PendingReview("請勿修改核對表欄名。")
    for index, expected in enumerate(rows, 2):
        if [sheet.cell(index, c).value for c in range(1, 12)] != expected:
            raise PendingReview(f"第 {index} 列來源欄位已更動；請在前段修正資料後重新核對。")
        if sheet.cell(index, 12).value != "Yes":
            raise PendingReview(f"第 {index} 列尚未確認，請完成 Data Review 的 Confirmed 欄。")
    name, reviewed, approved = (sign[c].value for c in ("B7", "B8", "B9"))
    if not isinstance(name, str) or not name.strip() or name.startswith("="):
        raise PendingReview("請在 Sign-off 填寫 Reviewer name。")
    if isinstance(reviewed, datetime):
        reviewed = reviewed.date()
    if isinstance(reviewed, str):
        try:
            reviewed = date.fromisoformat(reviewed)
        except ValueError:
            reviewed = None
    if not isinstance(reviewed, date):
        raise PendingReview("請填寫核對日期，格式 YYYY-MM-DD。")
    if approved != "Approved":
        raise PendingReview("請完成 Sign-off 的 Approval。")
    return {"reviewer": name.strip(), "review_date": reviewed.isoformat(),
            "source_id": digest, "row_count": len(rows),
            "review_file_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "accepted_at": datetime.now(timezone.utc).isoformat(),
            "record_type": "manual sign-off, identity not authenticated"}


def main():
    folder = Path(os.environ.get("CCIT_WORK_DIR", Path(__file__).resolve().parent)).resolve()
    rows = source_rows(folder)
    digest = fingerprint(folder, rows)
    path = folder / "CCIT_final_review.xlsx"
    if path.exists():
        old = load_workbook(path, read_only=True)
        old_digest = old["Sign-off"]["B12"].value
        old.close()
        if old_digest != digest:
            archive = folder / "review_history"
            archive.mkdir(exist_ok=True)
            shutil.copy2(path, archive / f"review_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.xlsx")
            create_review(path, rows, digest)
    else:
        create_review(path, rows, digest)
    print(f"最終人工核對表：{path}")
    signoff = verify_review(path, rows, digest)
    (folder / "final_signoff.json").write_text(
        json.dumps(signoff, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"人工核對與簽核完成：{len(rows)} 筆；核對人：{signoff['reviewer']}")


if __name__ == "__main__":
    try:
        main()
    except PendingReview as error:
        print(f"待人工核對／簽核：{error}")
        print("請填妥並關閉 Excel，再執行入口、選相同 PDF 與型態。尚未產生本次正式報表。")
        raise SystemExit(2)
