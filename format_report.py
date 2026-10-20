from pathlib import Path
import os
import csv
from openpyxl import Workbook

folder = Path(os.environ.get("CCIT_WORK_DIR", Path(__file__).resolve().parent)).resolve()

analysis_path = folder / "CCIT_v3_analysed.csv"

with analysis_path.open("r", newline="", encoding="utf-8-sig") as file:
    records = list(csv.DictReader(file))

workbook = Workbook()

analysis_sheet = workbook.active
analysis_sheet.title = "Sample Analysis"

settings_sheet = workbook.create_sheet("System Settings")

analysis_sheet.append([
    "Test No.",
    "Test ID",
    "Cycle 1 mbar",
    "Cycle 1 Pa",
    "Cycle 2 mbar",
    "Cycle 2 Pa",
    "Comments",
    "Instrument Result",
    "Sample Category",
    "Excluded from Reference",
    "WL Status",
])

category_labels = {
    "master_off": "Master MC-Off",
    "master_on": "Master MC-On",
    "positive": "Positive SPL",
    "sample": "Sample SPL",
}

wl_labels = {
    "不適用": "Not Applicable",
    "未達接近區間": "Below Near-WL Range",
    "接近 WL": "Near WL",
    "超過 WL": "Above WL",
}

for record in records:
    analysis_sheet.append([
        int(record["test_no"]),
        float(record["test_id"]),
        float(record["cycle1_mbar"]),
        float(record["cycle1_pa"]),
        float(record["cycle2_mbar"]),
        float(record["cycle2_pa"]),
        record["comments"],
        record["instrument_result"],
        category_labels[record["sample_category"]],
        "Yes" if record["initial_excluded"] == "True" else "No",
        wl_labels[record["wl_status"]],
    ])

for row in analysis_sheet.iter_rows(min_row=2, min_col=3, max_col=6):
    for cell in row:
        cell.number_format = "0.0"

print("工作表名稱：", workbook.sheetnames)
print("已填入資料筆數：", len(records))
print("樣品工作表總列數：", analysis_sheet.max_row)


from openpyxl.styles import Font, PatternFill, Alignment

header_fill = PatternFill("solid", fgColor="244062")
above_wl_fill = PatternFill("solid", fgColor="FCE4D6")
near_wl_fill = PatternFill("solid", fgColor="FFF2CC")
excluded_fill = PatternFill("solid", fgColor="F8CBAD")

for cell in analysis_sheet[1]:
    cell.fill = header_fill
    cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

analysis_sheet.row_dimensions[1].height = 36

for row in analysis_sheet.iter_rows(min_row=2):
    for cell in row:
        cell.font = Font(name="Calibri", size=11)
        cell.alignment = Alignment(vertical="center")

    status = row[10].value

    if status == "Above WL":
        for cell in row:
            cell.fill = above_wl_fill

    elif status == "Near WL":
        for cell in row:
            cell.fill = near_wl_fill

    if row[9].value == "Yes":
        row[9].fill = excluded_fill
        row[9].font = Font(name="Calibri", bold=True)

column_widths = {
    "A": 12,
    "B": 12,
    "C": 16,
    "D": 16,
    "E": 16,
    "F": 16,
    "G": 42,
    "H": 19,
    "I": 20,
    "J": 24,
    "K": 26,
}

for column, width in column_widths.items():
    analysis_sheet.column_dimensions[column].width = width

analysis_sheet.freeze_panes = "A2"
analysis_sheet.auto_filter.ref = analysis_sheet.dimensions

print("樣品表排版與顏色設定完成")


settings_path = folder / "CCIT_v3_system_settings.csv"

with settings_path.open("r", newline="", encoding="utf-8-sig") as file:
    settings_records = list(csv.DictReader(file))

print("系統參數筆數：", len(settings_records))

for record in settings_records:
    print(record)

settings_values = {}

for record in settings_records:
    settings_values[record["parameter"]] = float(record["value"])

product_type = settings_records[0]["product_type"]

settings_sheet.append(["Product Type", product_type.capitalize()])
settings_sheet.append([])

settings_sheet.append(["Parameter", "Minimum", "Maximum", "Value"])

for parameter in [
    "Cycle 1 mbar",
    "Cycle 1 Pa",
    "Cycle 2 mbar",
    "Cycle 2 Pa",
]:
    settings_sheet.append([
        parameter,
        settings_values[parameter + " min"],
        settings_values[parameter + " max"],
        None,
    ])

for parameter in ["System Vacuum", "Disable Sys Vac"]:
    settings_sheet.append([
        parameter,
        None,
        None,
        settings_values[parameter],
    ])

for row in settings_sheet.iter_rows():
    for cell in row:
        cell.font = Font(name="Calibri", size=11)
        cell.alignment = Alignment(vertical="center")

for cell in settings_sheet[3]:
    cell.fill = header_fill
    cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="center")

settings_sheet["A1"].font = Font(name="Calibri", bold=True)
settings_sheet.row_dimensions[3].height = 28

for row in settings_sheet.iter_rows(min_row=4, min_col=2, max_col=4):
    for cell in row:
        cell.number_format = "0.0"

settings_sheet.column_dimensions["A"].width = 26
settings_sheet.column_dimensions["B"].width = 16
settings_sheet.column_dimensions["C"].width = 16
settings_sheet.column_dimensions["D"].width = 16

settings_sheet.freeze_panes = "B4"

output_path = folder / "CCIT_v3_report.xlsx"
workbook.save(output_path)

print("Excel 報表已儲存：", output_path)
